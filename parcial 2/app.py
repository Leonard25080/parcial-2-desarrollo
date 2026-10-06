from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    send_from_directory,
    send_file
)

from werkzeug.security import check_password_hash
from werkzeug.utils import secure_filename

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from io import BytesIO
from pathlib import Path

import json


# =====================================================
# USUARIOS DEL SISTEMA
# =====================================================

USUARIOS = {

    # Usuario 1
    "Dulce_Hernandez":
        "scrypt:32768:8:1$XCywoPEa9cGbnKhN$d1b247c03aed16db143ad1e2606053ddd268cfb491cca20dc72291d7c91fbf95a552fca4a4d5902683ab9e2b5c698d02d3b7f4be4143987479065a64913f62e1",

    # Usuario 2
    "Alejandro_piña":
        "scrypt:32768:8:1$eWHwL1cebCJZTyXH$12cc3e80e9905b4f581a70cfac79e57585d3735c758ab8fdf15b7dd9757e2f36339fc42d16b120e1a689d8c883369ac0d93b1cc769f96485e0c013017c92b6ad"

}


# =====================================================
# CLAVE SECRETA DE FLASK
# =====================================================

SECRET_KEY = "75987d85e1fe5439a411d1fc04cfe85a67033cca531971cf963d01a676e1e737"


# =====================================================
# APLICACIÓN FLASK
# =====================================================

app = Flask(__name__)

app.secret_key = SECRET_KEY


# =====================================================
# CARPETAS Y ARCHIVOS
# =====================================================

BASE_DIR = Path(__file__).resolve().parent

FOTOS_DIR = BASE_DIR / "fotos"

CITAS_FILE = BASE_DIR / "citas.json"


FOTOS_DIR.mkdir(exist_ok=True)


# =====================================================
# EXTENSIONES DE IMAGEN PERMITIDAS
# =====================================================

EXTENSIONES_PERMITIDAS = {
    "png",
    "jpg",
    "jpeg",
    "webp"
}


# =====================================================
# COMPROBAR EXTENSIÓN DE ARCHIVO
# =====================================================

def extension_permitida(nombre_archivo):

    return (
        "." in nombre_archivo
        and nombre_archivo.rsplit(".", 1)[1].lower()
        in EXTENSIONES_PERMITIDAS
    )


# =====================================================
# CARGAR CITAS
# =====================================================

def cargar_citas():

    if not CITAS_FILE.exists():
        return []


    try:

        with open(
            CITAS_FILE,
            "r",
            encoding="utf-8"
        ) as archivo:

            return json.load(archivo)


    except (
        json.JSONDecodeError,
        OSError
    ):

        return []


# =====================================================
# GUARDAR CITAS
# =====================================================

def guardar_citas(citas):

    with open(
        CITAS_FILE,
        "w",
        encoding="utf-8"
    ) as archivo:

        json.dump(
            citas,
            archivo,
            ensure_ascii=False,
            indent=4
        )


# =====================================================
# MOSTRAR FOTOGRAFÍAS
# =====================================================

@app.route("/fotos/<nombre>")
def foto(nombre):

    if not session.get("logged_in"):

        return redirect(
            url_for("login")
        )


    return send_from_directory(
        FOTOS_DIR,
        nombre
    )


# =====================================================
# INICIO DE SESIÓN
# =====================================================

@app.route(
    "/login",
    methods=["GET", "POST"]
)
def login():

    if session.get("logged_in"):

        return redirect(
            url_for("index")
        )


    if request.method == "POST":

        usuario = request.form.get(
            "username",
            ""
        ).strip()


        password = request.form.get(
            "password",
            ""
        )


        # Buscar usuario

        password_hash = USUARIOS.get(
            usuario
        )


        # Comprobar contraseña

        if (
            password_hash
            and check_password_hash(
                password_hash,
                password
            )
        ):

            session["logged_in"] = True

            session["usuario"] = usuario


            return redirect(
                url_for("index")
            )


        return render_template(
            "login.html",
            error="Usuario o contraseña incorrectos"
        )


    return render_template(
        "login.html",
        error=None
    )


# =====================================================
# PÁGINA PRINCIPAL
# =====================================================

@app.route("/")
def index():

    if not session.get("logged_in"):

        return redirect(
            url_for("login")
        )


    citas = cargar_citas()


    # Ordenar citas por fecha y hora

    citas = sorted(
        citas,
        key=lambda cita: (
            cita.get("fecha", ""),
            cita.get("hora", "")
        )
    )


    return render_template(
        "index.html",
        usuario=session.get("usuario"),
        citas=citas
    )


# =====================================================
# REGISTRAR NUEVA CITA
# =====================================================

@app.route(
    "/nueva_cita",
    methods=["POST"]
)
def nueva_cita():

    if not session.get("logged_in"):

        return redirect(
            url_for("login")
        )


    citas = cargar_citas()


    # =================================================
    # DATOS DE LA MASCOTA
    # =================================================

    nombre = request.form.get(
        "nombre",
        ""
    ).strip()


    tipo = request.form.get(
        "tipo",
        ""
    ).strip()


    servicio = request.form.get(
        "servicio",
        ""
    ).strip()


    fecha = request.form.get(
        "fecha",
        ""
    ).strip()


    hora = request.form.get(
        "hora",
        ""
    ).strip()


    costo = request.form.get(
        "costo",
        ""
    ).strip()


    # =================================================
    # VALIDAR CAMPOS
    # =================================================

    if not all([
        nombre,
        tipo,
        servicio,
        fecha,
        hora,
        costo
    ]):

        return redirect(
            url_for("index")
        )


    # =================================================
    # VALIDAR COSTO
    # =================================================

    try:

        costo_numero = float(
            costo
        )


        if costo_numero < 0:

            return redirect(
                url_for("index")
            )


    except ValueError:

        return redirect(
            url_for("index")
        )


    # =================================================
    # FOTOGRAFÍA
    # =================================================

    archivo = request.files.get(
        "foto"
    )


    nombre_foto = ""


    if (
        archivo
        and archivo.filename
        and extension_permitida(
            archivo.filename
        )
    ):

        nombre_foto = secure_filename(
            archivo.filename
        )


        archivo.save(
            FOTOS_DIR / nombre_foto
        )


    # =================================================
    # CREAR CITA
    # =================================================

    nueva = {

        "nombre": nombre,

        "tipo": tipo,

        "servicio": servicio,

        "fecha": fecha,

        "hora": hora,

        "costo": costo_numero,

        "foto": nombre_foto,

        "registrado_por":
            session.get("usuario")
    }


    citas.append(
        nueva
    )


    guardar_citas(
        citas
    )


    return redirect(
        url_for("index")
    )


# =====================================================
# EXPORTAR CITAS A EXCEL
# =====================================================

@app.route("/exportar_excel")
def exportar_excel():

    # Comprobar que exista una sesión iniciada

    if not session.get("logged_in"):

        return redirect(
            url_for("login")
        )


    citas = cargar_citas()


    # =================================================
    # CREAR ARCHIVO DE EXCEL
    # =================================================

    libro = Workbook()

    hoja = libro.active

    hoja.title = "Citas"


    # =================================================
    # ENCABEZADOS
    # =================================================

    encabezados = [

        "Nombre de la mascota",

        "Tipo de animal",

        "Servicio",

        "Fecha",

        "Hora",

        "Costo",

        "Fotografía",

        "Registrado por"

    ]


    hoja.append(
        encabezados
    )


    # =================================================
    # FORMATO DE ENCABEZADOS
    # =================================================

    for celda in hoja[1]:

        celda.font = Font(
            bold=True
        )

        celda.alignment = Alignment(
            horizontal="center"
        )


    # =================================================
    # AGREGAR CITAS
    # =================================================

    for cita in citas:

        hoja.append([

            cita.get(
                "nombre",
                ""
            ),

            cita.get(
                "tipo",
                ""
            ),

            cita.get(
                "servicio",
                ""
            ),

            cita.get(
                "fecha",
                ""
            ),

            cita.get(
                "hora",
                ""
            ),

            cita.get(
                "costo",
                0
            ),

            cita.get(
                "foto",
                ""
            ),

            cita.get(
                "registrado_por",
                ""
            )

        ])


    # =================================================
    # FORMATO DE COSTO
    # =================================================

    for fila in range(
        2,
        hoja.max_row + 1
    ):

        hoja.cell(
            row=fila,
            column=6
        ).number_format = '$#,##0.00'


    # =================================================
    # AJUSTAR ANCHO DE COLUMNAS
    # =================================================

    hoja.column_dimensions["A"].width = 25

    hoja.column_dimensions["B"].width = 18

    hoja.column_dimensions["C"].width = 25

    hoja.column_dimensions["D"].width = 15

    hoja.column_dimensions["E"].width = 12

    hoja.column_dimensions["F"].width = 15

    hoja.column_dimensions["G"].width = 25

    hoja.column_dimensions["H"].width = 22


    # =================================================
    # CONGELAR ENCABEZADOS
    # =================================================

    hoja.freeze_panes = "A2"


    # =================================================
    # GUARDAR EXCEL EN MEMORIA
    # =================================================

    archivo_excel = BytesIO()


    libro.save(
        archivo_excel
    )


    archivo_excel.seek(
        0
    )


    # =================================================
    # DESCARGAR ARCHIVO
    # =================================================

    return send_file(

        archivo_excel,

        as_attachment=True,

        download_name="citas_veterinaria_tlacua.xlsx",

        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )

    )


# =====================================================
# CERRAR SESIÓN
# =====================================================

@app.route("/logout")
def logout():

    session.clear()


    return redirect(
        url_for("login")
    )


# =====================================================
# EJECUTAR FLASK
# =====================================================

if __name__ == "__main__":

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )